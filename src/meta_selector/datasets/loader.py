"""Local file readers: no download, target inference, encoding or preprocessing."""
import codecs
import csv
import re
from collections import Counter
from pathlib import Path

import pandas as pd
import yaml

from meta_selector.paths import ProjectPaths, dataset_defaults, project_paths, read_yaml_mapping
from .models import (Dataset, DatasetMetadata, DatasetReadError, DatasetStatus,
                     DatasetValidationError, FileMetadata, ProblemMetadata, SourceMetadata)


DATASET_ID_PATTERN = re.compile(r"^_?[a-z][a-z0-9]*(?:_[a-z0-9]+)*$")


def check_dataset_id(dataset_id: str) -> None:
    if not isinstance(dataset_id, str) or not DATASET_ID_PATTERN.fullmatch(dataset_id):
        raise DatasetReadError(DatasetStatus.INVALID_METADATA,
            "Некорректный dataset_id: используйте snake_case из строчных латинских букв и цифр, "
            "без пробелов и путей. Для тестовых наборов допустим начальный underscore.")


def _text(mapping: dict, key: str, context: str, *, allow_empty: bool = False) -> str:
    value = mapping.get(key)
    if not isinstance(value, str) or (not allow_empty and not value.strip()):
        raise ValueError(f"Обязательное поле {context}{key} должно быть строкой"
                         + ("" if allow_empty else " и не может быть пустым"))
    return value


def _section(mapping: dict, name: str) -> dict:
    section = mapping.get(name)
    if not isinstance(section, dict):
        raise ValueError(f"Обязательная секция {name} должна содержать mapping")
    return section


def read_metadata(dataset_id: str, paths: ProjectPaths) -> DatasetMetadata:
    check_dataset_id(dataset_id)
    defaults = dataset_defaults(paths)
    path = paths.raw_data / dataset_id / "metadata.yaml"
    try:
        if not path.is_file():
            raise ValueError(f"Не найден metadata.yaml: {path}")
        raw = read_yaml_mapping(path)
        declared = _text(raw, "dataset_id", "")
        if declared != dataset_id:
            raise ValueError(f"dataset_id {declared!r} не совпадает с именем папки {dataset_id!r}")
        name, target = _text(raw, "name", ""), _text(raw, "target", "")
        task_type = _text(raw, "task_type", "")
        if task_type != "classification":
            raise ValueError("task_type должен быть classification")
        source = _section(raw, "source")
        source_type = _text(source, "type", "source.")
        url = _text(source, "url", "source.", allow_empty=True)
        external_id = source.get("external_id")
        if isinstance(external_id, bool) or not isinstance(external_id, (str, int, type(None))):
            raise ValueError("source.external_id должен быть строкой, целым числом или null")
        file = _section(raw, "file")
        filename = _text(file, "name", "file.")
        if Path(filename).name != filename or filename in (".", "..") or "\\" in filename:
            raise ValueError("file.name должен быть именем файла внутри папки датасета, без пути")
        file_format = _text(file, "format", "file.")
        if file_format not in ("csv", "parquet", "xlsx"):
            raise ValueError("file.format должен быть csv, parquet или xlsx")
        if Path(filename).suffix.lower() != "." + file_format:
            raise ValueError("Расширение file.name не совпадает с file.format")
        separator = file.get("separator", defaults.get("csv", {}).get("separator", ","))
        encoding = file.get("encoding", defaults.get("csv", {}).get("encoding", "utf-8"))
        sheet = file.get("sheet_name", defaults.get("xlsx", {}).get("sheet_name"))
        if file_format == "csv":
            if not isinstance(separator, str) or len(separator) != 1 or separator in "\r\n":
                raise ValueError("file.separator должен быть одним символом-разделителем CSV")
            if not isinstance(encoding, str):
                raise ValueError("file.encoding должен быть названием кодировки CSV")
            codecs.lookup(encoding)
        if file_format == "xlsx" and (isinstance(sheet, bool) or not isinstance(sheet, (str, int, type(None)))
                                    or isinstance(sheet, int) and sheet < 0):
            raise ValueError("file.sheet_name должен быть именем листа, индексом >= 0 или null")
        problem_type = _text(_section(raw, "problem"), "type", "problem.")
        if problem_type not in ("binary", "multiclass"):
            raise ValueError("problem.type должен быть binary или multiclass")
        notes = raw.get("notes", "")
        if not isinstance(notes, str):
            raise ValueError("notes должен быть строкой")
        return DatasetMetadata(dataset_id, name, task_type, target,
            SourceMetadata(source_type, url, str(external_id) if external_id is not None else None),
            FileMetadata(filename, file_format, separator, encoding, sheet), ProblemMetadata(problem_type), notes)
    except (OSError, ValueError, TypeError, LookupError, yaml.YAMLError) as exc:
        raise DatasetReadError(DatasetStatus.INVALID_METADATA, str(exc)) from exc


def _check_headers(headers: list) -> None:
    duplicates = [name for name, count in Counter(headers).items() if count > 1]
    if duplicates:
        raise DatasetReadError(DatasetStatus.INVALID,
                               "Повторяющиеся исходные названия колонок: " + ", ".join(map(repr, duplicates)))


def read_frame(metadata: DatasetMetadata, paths: ProjectPaths) -> pd.DataFrame:
    folder = paths.raw_data / metadata.dataset_id
    path = folder / metadata.file.name
    if not path.is_file():
        raise DatasetReadError(DatasetStatus.MISSING_FILE, f"Не найден локальный файл: {path}")
    if not path.resolve().is_relative_to(folder.resolve()):
        raise DatasetReadError(DatasetStatus.INVALID_METADATA, "Файл должен находиться внутри папки датасета")
    options = metadata.file
    try:
        if options.format == "csv":
            with path.open(encoding=options.encoding, newline="") as stream:
                reader = csv.reader(stream, delimiter=options.separator)
                header = next((row for row in reader if row and not (len(row) == 1 and not row[0].strip())), [])
            if header:
                header[0] = header[0].removeprefix("\ufeff")
            _check_headers(header)
            frame = pd.read_csv(path, sep=options.separator, encoding=options.encoding)
        elif options.format == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet_key = 0 if options.sheet_name is None else options.sheet_name
                worksheet = workbook.worksheets[sheet_key] if isinstance(sheet_key, int) else workbook[sheet_key]
                header = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
                while header and header[-1] is None:
                    header.pop()
                _check_headers(header)
            finally:
                workbook.close()
            frame = pd.read_excel(path, sheet_name=sheet_key, engine="openpyxl")
        else:
            import pyarrow.parquet as pq

            _check_headers(pq.read_schema(path).names)
            frame = pd.read_parquet(path, engine="pyarrow")
        _check_headers(frame.columns.tolist())
        return frame
    except DatasetReadError:
        raise
    except Exception as exc:
        raise DatasetReadError(DatasetStatus.INVALID,
                               f"Не удалось прочитать {metadata.file.name}: {type(exc).__name__}: {exc}") from exc


def load_dataset(dataset_id: str, *, project_root: str | Path | None = None) -> Dataset:
    from .validator import _inspect_dataset

    report, frame = _inspect_dataset(dataset_id, project_paths(project_root))
    if not report.is_valid:
        raise DatasetValidationError(report)
    metadata = report.metadata
    assert metadata is not None and frame is not None
    return Dataset(dataset_id, frame.drop(columns=[metadata.target]), frame[metadata.target], metadata)
